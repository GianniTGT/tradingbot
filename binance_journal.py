"""
Binance Trade Journal -- holt echte Trade-History und exportiert als Excel
Strategie: EMA20 Pullback | Long only | 1% Risiko | 2:1 RR
"""

import json, time, hmac, hashlib, requests
import pandas as pd
from datetime import datetime
from pathlib import Path
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Config ─────────────────────────────────────────────────────────────────────
import os as _os
API_KEY    = _os.environ.get("BINANCE_API_KEY")
API_SECRET = _os.environ.get("BINANCE_API_SECRET")

if not API_KEY or not API_SECRET:
    with open(_os.path.join(_os.path.dirname(__file__), "binance_config.json")) as f:
        cfg = json.load(f)
    API_KEY    = cfg["api_key"]
    API_SECRET = cfg["api_secret"]

SYMBOLS = [
    "BTCUSDT","ETHUSDT","SOLUSDT","BNBUSDT","LINKUSDT",
    "NEARUSDT","AVAXUSDT","MAGICUSDT","DOTUSDT","ADAUSDT",
    "XRPUSDT","SUIUSDT","INJUSDT","APTUSDT","ARBUSDT"
]

SPOT_BASE    = "https://api.binance.com"
FUTURES_BASE = "https://fapi.binance.com"

# ── Helper ─────────────────────────────────────────────────────────────────────
def sign(params: dict) -> dict:
    params["timestamp"] = int(time.time() * 1000)
    query = "&".join(f"{k}={v}" for k, v in params.items())
    params["signature"] = hmac.new(
        API_SECRET.encode(), query.encode(), hashlib.sha256
    ).hexdigest()
    return params

def get_trades(base_url, endpoint, symbol):
    params = sign({"symbol": symbol, "limit": 1000})
    r = requests.get(
        f"{base_url}{endpoint}",
        params=params,
        headers={"X-MBX-APIKEY": API_KEY},
        timeout=10
    )
    return r.json() if r.status_code == 200 else []

# ── Trades holen ───────────────────────────────────────────────────────────────
print("Hole Spot-Trades...")
spot_rows = []
for sym in SYMBOLS:
    trades = get_trades(SPOT_BASE, "/api/v3/myTrades", sym)
    for t in trades:
        spot_rows.append({
            "Symbol":    sym.replace("USDT", ""),
            "Side":      "BUY" if t["isBuyer"] else "SELL",
            "Datum":     datetime.fromtimestamp(t["time"] / 1000).strftime("%Y-%m-%d %H:%M"),
            "Preis":     float(t["price"]),
            "Menge":     float(t["qty"]),
            "Wert_USD":  float(t["quoteQty"]),
            "Gebuehr":   float(t["commission"]),
            "Order_ID":  t["orderId"],
        })

print("Hole Futures-Trades...")
fut_rows = []
for sym in SYMBOLS:
    trades = get_trades(FUTURES_BASE, "/fapi/v1/userTrades", sym)
    for t in trades:
        fut_rows.append({
            "Symbol":    sym.replace("USDT", ""),
            "Side":      t.get("side", "?"),
            "Datum":     datetime.fromtimestamp(t["time"] / 1000).strftime("%Y-%m-%d %H:%M"),
            "Preis":     float(t["price"]),
            "Menge":     float(t["qty"]),
            "Wert_USD":  float(t["quoteQty"]),
            "RealPnL":   float(t.get("realizedPnl", 0)),
            "Gebuehr":   float(t["commission"]),
            "Order_ID":  t["orderId"],
        })

print(f"  Spot: {len(spot_rows)} Fills | Futures: {len(fut_rows)} Fills")

# ── Orders aggregieren (Teilausführungen zusammenfassen) ──────────────────────
def aggregate_orders(rows):
    """Gruppiert einzelne Fills nach Order-ID zu einer Order."""
    if not rows:
        return pd.DataFrame()
    df = pd.DataFrame(rows)
    orders = []
    for order_id, grp in df.groupby("Order_ID"):
        avg_price = (grp["Preis"] * grp["Menge"]).sum() / grp["Menge"].sum()
        orders.append({
            "Symbol":   grp["Symbol"].iloc[0],
            "Side":     grp["Side"].iloc[0],
            "Datum":    grp["Datum"].min(),
            "Preis":    round(avg_price, 6),
            "Menge":    round(grp["Menge"].sum(), 8),
            "Wert_USD": round(grp["Wert_USD"].sum(), 4),
            "Gebuehr":  round(grp["Gebuehr"].sum(), 6),
            "Order_ID": order_id,
        })
    return pd.DataFrame(orders).sort_values("Datum").reset_index(drop=True)

spot_orders = aggregate_orders(spot_rows)
print(f"  Aggregiert: {len(spot_orders)} Orders (Spot)")

# ── Trades paaren (BUY -> SELL) ───────────────────────────────────────────────
def pair_trades(orders_df, risk_pct=1.0):
    """Paart BUY/SELL Orders zu vollständigen Trades."""
    if orders_df.empty:
        return pd.DataFrame()

    journal = []
    for sym in orders_df["Symbol"].unique():
        s = orders_df[orders_df["Symbol"] == sym].copy().reset_index(drop=True)
        buys  = s[s["Side"].isin(["BUY", "LONG"])].copy().reset_index(drop=True)
        sells = s[s["Side"].isin(["SELL", "SHORT"])].copy().reset_index(drop=True)

        used = set()
        for _, buy in buys.iterrows():
            # nächsten Sell nach diesem Buy
            candidates = sells[
                (sells["Datum"] > buy["Datum"]) &
                (~sells.index.isin(used))
            ]

            if candidates.empty:
                journal.append({
                    "Nr":           "",
                    "Symbol":       sym,
                    "Datum Entry":  buy["Datum"],
                    "Datum Exit":   "offen",
                    "Entry Preis":  buy["Preis"],
                    "Exit Preis":   "-",
                    "Menge":        buy["Menge"],
                    "Invest USD":   round(buy["Wert_USD"], 2),
                    "Gewinn/Verlust USD": "-",
                    "Gewinn/Verlust Netto": "-",
                    "PnL %":        "-",
                    "RR":           "-",
                    "Ergebnis":     "offen",
                    "Gebuehr":      round(buy["Gebuehr"], 4),
                    "Dauer":        "-",
                })
                continue

            sell = candidates.iloc[0]
            used.add(sell.name)

            qty      = min(buy["Menge"], sell["Menge"])
            entry    = buy["Preis"]
            exit_p   = sell["Preis"]
            invest   = round(entry * qty, 2)
            pnl_usd  = round((exit_p - entry) * qty, 2)
            gebuehr  = round(buy["Gebuehr"] + sell["Gebuehr"], 4)
            pnl_net  = round(pnl_usd - gebuehr, 2)
            pnl_pct  = round((exit_p - entry) / entry * 100, 2)

            # RR: Gewinn/Verlust relativ zum definierten Risiko (1% des Investments)
            risk_usd = round(invest * risk_pct / 100, 2)
            rr       = round(pnl_usd / risk_usd, 2) if risk_usd != 0 else 0

            # Handelsdauer
            try:
                d1 = datetime.strptime(buy["Datum"], "%Y-%m-%d %H:%M")
                d2 = datetime.strptime(sell["Datum"], "%Y-%m-%d %H:%M")
                delta = d2 - d1
                h, m = divmod(int(delta.total_seconds() / 60), 60)
                dauer = f"{h}h {m}m"
            except Exception:
                dauer = "-"

            journal.append({
                "Nr":           "",
                "Symbol":       sym,
                "Datum Entry":  buy["Datum"],
                "Datum Exit":   sell["Datum"],
                "Entry Preis":  entry,
                "Exit Preis":   exit_p,
                "Menge":        qty,
                "Invest USD":   invest,
                "Gewinn/Verlust USD":   pnl_usd,
                "Gewinn/Verlust Netto": pnl_net,
                "PnL %":        pnl_pct,
                "RR":           rr,
                "Ergebnis":     "WIN" if pnl_usd > 0 else "LOSS",
                "Gebuehr":      gebuehr,
                "Dauer":        dauer,
            })

    df = pd.DataFrame(journal)
    # Nummerierung
    df["Nr"] = range(1, len(df) + 1)
    return df

spot_journal = pair_trades(spot_orders)
closed = spot_journal[spot_journal["Ergebnis"].isin(["WIN","LOSS"])].copy()

# ── Statistik berechnen ────────────────────────────────────────────────────────
def calc_stats(df):
    closed = df[df["Ergebnis"].isin(["WIN","LOSS"])].copy()
    if closed.empty:
        return {}
    wins   = closed[closed["Ergebnis"] == "WIN"]
    losses = closed[closed["Ergebnis"] == "LOSS"]

    total_gewinn = round(wins["Gewinn/Verlust USD"].sum(), 2)
    total_verlust = round(losses["Gewinn/Verlust USD"].sum(), 2)
    net_pnl = round(closed["Gewinn/Verlust Netto"].sum(), 2)
    winrate = round(len(wins) / len(closed) * 100, 1)
    avg_win = round(wins["Gewinn/Verlust USD"].mean(), 2) if not wins.empty else 0
    avg_loss = round(losses["Gewinn/Verlust USD"].mean(), 2) if not losses.empty else 0
    pf = round(total_gewinn / abs(total_verlust), 2) if total_verlust != 0 else float("inf")
    avg_rr = round(closed["RR"].mean(), 2)
    best = round(wins["Gewinn/Verlust USD"].max(), 2) if not wins.empty else 0
    worst = round(losses["Gewinn/Verlust USD"].min(), 2) if not losses.empty else 0

    return {
        "Trades gesamt":         len(closed),
        "Offene Positionen":     len(df[df["Ergebnis"] == "offen"]),
        "Wins":                  len(wins),
        "Losses":                len(losses),
        "Winrate":               f"{winrate}%",
        "Gesamt Gewinn":         f"+${total_gewinn}",
        "Gesamt Verlust":        f"${total_verlust}",
        "Net PnL (nach Gebühr)": f"${net_pnl}",
        "Profit Factor":         pf,
        "Avg RR":                avg_rr,
        "Avg Gewinn/Trade":      f"+${avg_win}",
        "Avg Verlust/Trade":     f"${avg_loss}",
        "Bester Trade":          f"+${best}",
        "Schlechtester Trade":   f"${worst}",
    }

stats = calc_stats(spot_journal)
print("\n--- ZUSAMMENFASSUNG ---")
for k, v in stats.items():
    print(f"  {k}: {v}")

# ── Excel Export mit Formatierung ─────────────────────────────────────────────
output = Path("trade_journal.xlsx")

# Farben
GREEN_BG  = PatternFill("solid", fgColor="C6EFCE")
RED_BG    = PatternFill("solid", fgColor="FFC7CE")
OPEN_BG   = PatternFill("solid", fgColor="FFEB9C")
HEADER_BG = PatternFill("solid", fgColor="1F4E79")
TOTAL_BG  = PatternFill("solid", fgColor="2E75B6")
GREEN_FT  = Font(bold=True, color="276221")
RED_FT    = Font(bold=True, color="9C0006")
WHITE_FT  = Font(bold=True, color="FFFFFF")
BOLD      = Font(bold=True)
CENTER    = Alignment(horizontal="center", vertical="center")
thin = Side(style="thin", color="CCCCCC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEADER_BG
        cell.font = WHITE_FT
        cell.alignment = CENTER
        cell.border = BORDER

def autofit(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 35)

with pd.ExcelWriter(output, engine="openpyxl") as writer:

    # ── Tab 1: Journal (Hauptansicht) ─────────────────────────────────────────
    cols_order = ["Nr","Symbol","Datum Entry","Datum Exit","Dauer",
                  "Entry Preis","Exit Preis","Menge","Invest USD",
                  "Gewinn/Verlust USD","Gewinn/Verlust Netto","PnL %","RR",
                  "Ergebnis","Gebuehr"]

    journal_sorted = spot_journal[cols_order].sort_values("Nr", ascending=False)
    journal_sorted.to_excel(writer, sheet_name="Trade Journal", index=False)

    ws = writer.sheets["Trade Journal"]
    style_header(ws)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        ergebnis = row[cols_order.index("Ergebnis")].value
        for cell in row:
            cell.border = BORDER
            cell.alignment = CENTER
        if ergebnis == "WIN":
            for cell in row:
                cell.fill = GREEN_BG
            row[cols_order.index("Ergebnis")].font = GREEN_FT
            row[cols_order.index("Gewinn/Verlust USD")].font = GREEN_FT
        elif ergebnis == "LOSS":
            for cell in row:
                cell.fill = RED_BG
            row[cols_order.index("Ergebnis")].font = RED_FT
            row[cols_order.index("Gewinn/Verlust USD")].font = RED_FT
        elif ergebnis == "offen":
            for cell in row:
                cell.fill = OPEN_BG

    # Summen-Zeile
    last = ws.max_row + 1
    closed_df = journal_sorted[journal_sorted["Ergebnis"].isin(["WIN","LOSS"])]
    ws.cell(last, 1, "TOTAL")
    ws.cell(last, cols_order.index("Invest USD") + 1, round(closed_df["Invest USD"].sum(), 2))
    ws.cell(last, cols_order.index("Gewinn/Verlust USD") + 1, round(closed_df["Gewinn/Verlust USD"].sum(), 2))
    ws.cell(last, cols_order.index("Gewinn/Verlust Netto") + 1, round(closed_df["Gewinn/Verlust Netto"].sum(), 2))
    ws.cell(last, cols_order.index("Gebuehr") + 1, round(closed_df["Gebuehr"].sum(), 4))
    for cell in ws[last]:
        cell.fill = TOTAL_BG
        cell.font = WHITE_FT
        cell.alignment = CENTER
        cell.border = BORDER

    autofit(ws)
    ws.freeze_panes = "A2"

    # ── Tab 2: Zusammenfassung ────────────────────────────────────────────────
    stat_df = pd.DataFrame(list(stats.items()), columns=["Kennzahl", "Wert"])
    stat_df.to_excel(writer, sheet_name="Zusammenfassung", index=False)
    ws2 = writer.sheets["Zusammenfassung"]
    style_header(ws2)
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            cell.alignment = CENTER
        val = str(row[1].value or "")
        if val.startswith("+"):
            row[1].font = GREEN_FT
            row[1].fill = GREEN_BG
        elif val.startswith("-") or val.startswith("$-"):
            row[1].font = RED_FT
            row[1].fill = RED_BG
        row[0].font = BOLD
    autofit(ws2)

    # ── Tab 3: Per-Coin Stats ─────────────────────────────────────────────────
    coin_rows = []
    for sym in closed["Symbol"].unique():
        s = closed[closed["Symbol"] == sym]
        wins = s[s["Ergebnis"] == "WIN"]
        losses = s[s["Ergebnis"] == "LOSS"]
        g = round(wins["Gewinn/Verlust USD"].sum(), 2)
        v = round(losses["Gewinn/Verlust USD"].sum(), 2)
        coin_rows.append({
            "Coin":         sym,
            "Trades":       len(s),
            "Wins":         len(wins),
            "Losses":       len(losses),
            "Winrate":      f"{round(len(wins)/len(s)*100,1)}%" if len(s) > 0 else "-",
            "Gewinn USD":   f"+${g}" if g >= 0 else f"${g}",
            "Verlust USD":  f"${v}",
            "Net PnL":      f"${round(g+v, 2)}",
            "Avg RR":       round(s["RR"].mean(), 2),
        })
    coin_df = pd.DataFrame(coin_rows)
    coin_df.to_excel(writer, sheet_name="Per-Coin Stats", index=False)
    ws3 = writer.sheets["Per-Coin Stats"]
    style_header(ws3)
    for row in ws3.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            cell.alignment = CENTER
        net = str(row[7].value or "")
        if not net.startswith("$-") and net != "$0":
            row[7].font = GREEN_FT
            row[7].fill = GREEN_BG
        else:
            row[7].font = RED_FT
            row[7].fill = RED_BG
    autofit(ws3)

    # ── Tab 4: Alle Rohdaten ──────────────────────────────────────────────────
    raw_df = pd.DataFrame(spot_rows).sort_values("Datum", ascending=False)
    raw_df.to_excel(writer, sheet_name="Rohdaten", index=False)
    ws4 = writer.sheets["Rohdaten"]
    style_header(ws4)
    autofit(ws4)

print(f"\nExcel gespeichert: {output.resolve()}")
